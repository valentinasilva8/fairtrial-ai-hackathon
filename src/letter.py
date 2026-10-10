"""Draft a submission to the UN Special Rapporteur on freedom of expression, with every sentence sourced.

1. build_sources() makes a numbered list of sources from the project's data:
   the case facts and timeline, the stress-test result, official pledges, and
   paragraphs from similar TrialWatch reports (with page links), and the sentences
   those reports state about the impact on their defendants (unconfirmed until a
   person fills verified_by).
2. A draft is a list of sections, each a list of sentences, and every sentence
   names the source ids it relies on. Gemini writes it (generate_draft), or
   template_draft() builds a plain version from the same sources without an LLM.
3. validate() rejects any sentence with no source, an unknown source, or a
   quotation that isn't in the cited source. Nothing is used until a person approves.
"""

from __future__ import annotations

import re
from dataclasses import dataclass

from src.argument_bank import CATEGORY_LABELS, page_link, same_drafting
from src.llm import LLMError, generate_json

RECIPIENT = "Special Rapporteur on the promotion and protection of the right to freedom of opinion and expression"
SECTIONS = [
    "Summary",
    "The case",
    "The reform and how it applies",
    "International standards and TrialWatch's findings in similar cases",
    "Requested action",
]


class SensitiveCaseError(ValueError):
    """Sensitive cases are never used in briefs."""


@dataclass
class Source:
    id: str
    label: str
    text: str
    url: str = ""
    verified: bool = True
    country: str = ""


def _yes(v) -> bool:
    return str(v).strip().lower() in {"true", "1", "yes"}


def build_sources(case: dict, events: list[dict], stress: dict | None, promises: list[dict],
                  past: list[dict]) -> list[Source]:
    """past: [{"case", "report_url", "pdf_url", "outcome", "outcome_confirmed", "arguments": {cat: [para]}}].

    `case` may be a seeded Indonesian case or a new one a lawyer describes (with "country" and "law");
    the stress test and pledges are optional because they exist only for a reformed law.
    """
    if _yes(case.get("sensitive")):
        raise SensitiveCaseError("This case is marked sensitive and can't be used in a brief.")
    cid = case["case_id"]
    country = case.get("country") or "Indonesia"
    law = case.get("law") or "the ITE (EIT) Law"
    src = [Source(
        f"case:{cid}", f"Case record: {case['name']}",
        f"{case['name']} ({case['role']}), prosecuted in {country} under {law}. "
        f"Complainant: {case.get('complainant', 'not given')} ({case.get('complainant_type', 'unknown')}). "
        f"Charged under: {case.get('article', 'not given')}. Reported: {case.get('year_reported', 'not given')}. "
        f"Status: {case.get('outcome', 'not given')} — {case.get('outcome_detail', '')}",
        case.get("source", ""), _yes(case.get("verified")), country=country,
    )]
    for i, e in enumerate(events, 1):
        src.append(Source(f"event:{cid}:{i}", f"Timeline, {e['date']}", f"{e['date']}: {e['description']}",
                          e.get("source", ""), _yes(e.get("verified"))))
    if stress:
        src.append(Source(
            f"stress:{cid}", "Precedent & Practice reform stress test (for lawyer review)",
            f"Verdict: {stress['verdict']}. " + " ".join(stress["reasons"]),
            "docs/stress_test_rules.md",
        ))
    for p in promises:
        src.append(Source(f"promise:{p['promise_id']}", f"Official pledge: {p['made_by']}",
                          f"{p['promise_text']} ({p['made_by']}, {p['date']}). Status: {p['status']}.",
                          p.get("source", ""), _yes(p.get("verified"))))
    tw_sources: list[Source] = []
    for m in past:
        slug = m["report_url"].rstrip("/").rsplit("/", 1)[-1]
        outcome = f"{m['outcome']} ({'confirmed' if m['outcome_confirmed'] else 'unconfirmed'})"
        if m.get("outcome_history"):
            outcome += "; history: " + "; ".join(m["outcome_history"])
        by = f" by {m['author']}" if m.get("author") else ""
        for cat, paras in m["arguments"].items():
            for para in paras:
                # The same paragraph reused in several reports is one source, naming every report.
                twin = next((t for t in tw_sources if same_drafting(t.text, para["text"])), None)
                if twin:
                    if m["case"] not in twin.label:
                        twin.label += f"; same paragraph in the report on {m['case']}{by}, p. {para['page']}"
                    continue
                twin = Source(
                    f"tw:{slug}:p{para['page']}",
                    f"TrialWatch fairness report on {m['case']}{by}, p. {para['page']} "
                    f"({CATEGORY_LABELS[cat]}; outcome: {outcome})",
                    para["text"], page_link(m.get("pdf_url", ""), para["page"]) or m["report_url"],
                )
                tw_sources.append(twin)
                src.append(twin)
        for imp in m.get("impacts", []):
            status = "confirmed" if imp["confirmed"] else "unconfirmed"
            src.append(Source(
                f"impact:{slug}:{imp['category']}",
                f"TrialWatch report, {m['case']}, p. {imp['page']} (impact on the defendant: {imp['label']}; {status})",
                imp["sentence"], imp["url"] or m["report_url"], imp["confirmed"],
            ))
    # de-duplicate ids (one paragraph can serve several categories)
    seen, out = set(), []
    for s in src:
        if s.id not in seen:
            seen.add(s.id)
            out.append(s)
    return out


QUOTE = re.compile(r"[“\"]([^”\"]{12,})[”\"]")


def _norm(s: str) -> str:
    """Words only, in order: ignores case, punctuation, brackets and spacing (e.g. "[T[he" -> "the")."""
    s = re.sub(r"[\[\]]", "", s.lower())
    return " ".join(re.findall(r"[a-z0-9]+", s))


def validate(draft: dict, sources: list[Source]) -> list[str]:
    """Problems with a draft; an empty list means every sentence is sourced."""
    by_id = {s.id: s for s in sources}
    errors = []
    sections = draft.get("sections", [])
    if not sections:
        return ["The draft has no sections."]
    for sec in sections:
        for n, sent in enumerate(sec.get("sentences", []), 1):
            where = f"{sec.get('heading', '?')} sentence {n}"
            ids = sent.get("sources", [])
            if not str(sent.get("text", "")).strip():
                errors.append(f"{where}: empty sentence.")
                continue
            if not ids:
                errors.append(f"{where}: no source cited.")
            unknown = [i for i in ids if i not in by_id]
            if unknown:
                errors.append(f"{where}: unknown source {', '.join(unknown)}.")
            for q in QUOTE.findall(sent["text"]):
                if not any(_norm(q) in _norm(by_id[i].text) for i in ids if i in by_id):
                    errors.append(f"{where}: quotation not found in the cited source: “{q[:60]}”.")
    return errors


def template_draft(sources: list[Source]) -> dict:
    """A plain, fully sourced draft built without an LLM."""
    case = sources[0]
    country = case.country or "Indonesia"
    stress = next((s for s in sources if s.id.startswith("stress:")), None)
    events = [s for s in sources if s.id.startswith("event:")]
    promises = [s for s in sources if s.id.startswith("promise:")]
    tw = [s for s in sources if s.id.startswith("tw:")]
    impacts = [s for s in sources if s.id.startswith("impact:")]

    def label(s: Source) -> str:
        # A straight quote in a case title ('Katanyu "Pan"') would be read by validate() as the start of a quotation.
        return s.label.replace('"', "'")

    summary = [{"text": f"We write regarding the case of {case.label.split(': ', 1)[1]}.", "sources": [case.id]}]
    if stress:
        summary.append({"text": f"Our rule check of {country}'s reformed law gives this result: {stress.text}",
                        "sources": [stress.id]})
    sections = [
        {"heading": "Summary", "sentences": summary},
        {"heading": "The case", "sentences": [{"text": case.text, "sources": [case.id]}]
            + [{"text": e.text, "sources": [e.id]} for e in events]},
    ]
    if promises:
        sections.append({"heading": "The reform and how it applies", "sentences":
                         [{"text": p.text, "sources": [p.id]} for p in promises]})
    if stress:
        ask = (f"We ask that you consider raising this case with the Government of {country}, "
               "including whether the prosecution is consistent with the 2025 Constitutional Court ruling.")
        ask_sources = [stress.id] + [p.id for p in promises[:2]]
    else:
        ask = (f"We ask that you consider raising this case with the Government of {country}, including whether "
               "the prosecution is consistent with Article 19 of the ICCPR as applied in the reports cited above.")
        ask_sources = [case.id] + [s.id for s in tw[:2]]
    return {"sections": sections + [
        {"heading": "International standards and TrialWatch's findings in similar cases", "sentences":
            [{"text": f"{label(s)}: “{s.text[:300].rsplit(' ', 1)[0]}”", "sources": [s.id]} for s in tw]
            + [{"text": f"{label(s)}: “{s.text}”", "sources": [s.id]} for s in impacts]},
        {"heading": "Requested action", "sentences": [{"text": ask, "sources": ask_sources}]},
    ]}


DRAFT_SCHEMA = {
    "type": "object",
    "properties": {"sections": {"type": "array", "items": {
        "type": "object",
        "properties": {
            "heading": {"type": "string", "enum": SECTIONS},
            "sentences": {"type": "array", "items": {
                "type": "object",
                "properties": {"text": {"type": "string"}, "sources": {"type": "array", "items": {"type": "string"}}},
                "required": ["text", "sources"],
            }},
        },
        "required": ["heading", "sentences"],
    }}},
    "required": ["sections"],
}

SYSTEM = f"""You draft a submission to the UN {RECIPIENT} for a human rights organisation's lawyers to review.

Rules:
- Use ONLY the numbered sources provided. Do not add facts, dates, names, numbers or legal claims that are not in them.
- Every sentence must list the ids of the sources it relies on in "sources".
- If you quote, copy the words exactly from the cited source and put them in quotation marks.
- Sections, in order: {", ".join(SECTIONS)}.
- In the international-standards section, explain how the arguments made in TrialWatch fairness reports on similar past
  cases (legality, vagueness, broadness, necessity, proportionality) apply to this case,
  naming the past case and the report's author (given in the source label) for each point, e.g. "the TrialWatch
  fairness report by Lisa Davis on Poland v. Podlesna argued…". Do not write "TrialWatch argued": the analysis is
  the named author's. If a source label says the same paragraph appears in several reports, treat it as one point. Say an outcome "followed", never that an argument "caused" it.
  When you mention a past case's outcome, say whether it is confirmed, e.g. "(outcome not yet confirmed)";
  the case itself is not "unconfirmed".
- Sources whose id starts with "impact:" are sentences from a past TrialWatch report about what the prosecution did
  to that report's defendant (conviction and sentence, detention, mistreatment and so on). Cite them only for that
  past case, never as facts about the current case. Keep any number exactly as written, say it is unconfirmed when
  the source says so, and say the impact "followed" the prosecution; never that an argument caused anything.
- Describe the stress test, if there is one, as the authors' preliminary analysis for lawyer review, not a legal
  conclusion. If there are no stress-test or pledge sources, leave out "The reform and how it applies".
- Formal, factual tone. About 350 to 550 words."""


def _prompt(sources: list[Source], feedback: list[str]) -> str:
    lines = ["Sources:"]
    for s in sources:
        flag = "" if s.verified else " [unverified]"
        lines.append(f"[{s.id}] {s.label}{flag}\n{s.text}")
    if feedback:
        lines.append("\nYour previous draft was rejected for these reasons. Fix every one:\n- " + "\n- ".join(feedback))
    return "\n\n".join(lines)


def generate_draft(sources: list[Source], client=None, attempts: int = 3) -> tuple[dict, list[str]]:
    """Ask Gemini for a draft and re-ask with the problems until it validates. Returns (draft, problems)."""
    feedback: list[str] = []
    draft: dict = {}
    for _ in range(attempts):
        draft = generate_json(SYSTEM, _prompt(sources, feedback), DRAFT_SCHEMA, client=client, temperature=0.2)
        feedback = validate(draft, sources)
        if not feedback:
            return draft, []
    return draft, feedback


SUPPORT_SCHEMA = {
    "type": "object",
    "properties": {"checks": {"type": "array", "items": {
        "type": "object",
        "properties": {
            "n": {"type": "integer"},
            "verdict": {"type": "string", "enum": ["supported", "partly supported", "not supported"]},
            "evidence": {"type": "string"},
            "reason": {"type": "string"},
        },
        "required": ["n", "verdict", "evidence", "reason"],
    }}},
    "required": ["checks"],
}

SUPPORT_SYSTEM = """You check a draft legal letter sentence by sentence. For each numbered sentence, decide
whether its cited sources state what the sentence says: "supported" (every claim is in the sources),
"partly supported" (some claims, or an inference beyond the sources), or "not supported".
Each source starts with a label (for TrialWatch reports: the case, page, argument category and the case outcome);
the label counts as part of the source. Sentences in the "Requested action" section are requests: check only that
the facts they rest on are in the sources.
Give "evidence": the exact words from the cited source that support it (empty if none), and a one-line reason.
Judge only against the cited sources, not your own knowledge."""


def sentences_of(draft: dict) -> list[dict]:
    """Flat list of the draft's sentences with their section."""
    return [{**sent, "heading": sec["heading"]} for sec in draft["sections"] for sent in sec["sentences"]]


def check_support(draft: dict, sources: list[Source], client=None) -> list[dict]:
    """Second pass: does each sentence's cited source actually say it? Evidence is checked verbatim."""
    by_id = {s.id: s for s in sources}
    sents = sentences_of(draft)
    blocks = []
    for n, sent in enumerate(sents, 1):
        cited = "\n".join(f"[{i}] {by_id[i].label}: {by_id[i].text}" for i in sent["sources"] if i in by_id)
        blocks.append(f"Sentence {n}: {sent['text']}\nCited sources:\n{cited}")
    result = generate_json(SUPPORT_SYSTEM, "\n\n".join(blocks), SUPPORT_SCHEMA, client=client)
    checks = {c["n"]: c for c in result.get("checks", [])}
    out = []
    for n, sent in enumerate(sents, 1):
        c = checks.get(n, {"verdict": "not checked", "evidence": "", "reason": "No answer for this sentence."})
        evidence = c.get("evidence", "")
        # The checker's evidence must itself be real text from a cited source.
        cited_text = " ".join(_norm(f"{by_id[i].label}: {by_id[i].text}") for i in sent["sources"] if i in by_id)
        pieces = [_norm(p) for p in re.split(r"\.\.\.|…|\[\.\.\.\]", evidence) if _norm(p)]
        if c["verdict"] == "supported" and pieces and not all(p in cited_text for p in pieces):
            c = {**c, "verdict": "partly supported", "reason": "Quoted evidence not found in the source. " + c["reason"]}
        out.append({"n": n, "text": sent["text"], "heading": sent["heading"], **c})
    return out


def numbered_sources(draft: dict, sources: list[Source]) -> list[Source]:
    """Sources in the order the letter cites them: [1] is the first, [2] the second, …"""
    by_id = {s.id: s for s in sources}
    order: list[str] = []
    for sec in draft["sections"]:
        for sent in sec["sentences"]:
            for i in sent["sources"]:
                if i not in order and i in by_id:
                    order.append(i)
    return [by_id[i] for i in order]


CITE = re.compile(r"\[(\d+)\]")


def _plain_sentences(body: str) -> list[str]:
    lines = [p.strip() for p in body.split("\n")
             if p.strip() and not p.lstrip().startswith(("#", "**To:**", "**Re:**", "*Draft"))]
    return [s.strip() for p in lines for s in re.split(r"(?<=[.!?])\s+(?=[A-Z“\"])", p) if s.strip()]


def check_edited(text: str, numbered: list[Source], original: str = "") -> list[str]:
    """Problems in a letter a reviewer has edited by hand (no AI): broken citations, changed quotations,
    or new sentences without a citation (sentences already in the original draft are not re-flagged)."""
    body = text.split("### Sources")[0]
    problems = []
    bad = sorted({int(n) for n in CITE.findall(body) if not 1 <= int(n) <= len(numbered)})
    if bad:
        problems.append(f"Citation numbers with no matching source: {', '.join(f'[{n}]' for n in bad)}.")
    for q in QUOTE.findall(body):
        if not any(_norm(q) in _norm(s.text) for s in numbered):
            problems.append(f"Quotation not found word for word in any source: “{q[:70]}”.")
    before = set(_plain_sentences(original.split("### Sources")[0])) if original else set()
    uncited = [s for s in _plain_sentences(body) if len(s) > 40 and not CITE.search(s) and s not in before]
    if uncited:
        problems.append(f"{len(uncited)} sentence(s) have no citation, e.g. “{uncited[0][:70]}”.")
    return problems


def parse_edited(text: str, numbered: list[Source]) -> dict:
    """Turn an edited letter back into sections of sentences with their cited source ids,
    so the same sentence-by-sentence support check can run on the edited text."""
    body = text.split("### Sources")[0]
    sections, heading = [], "Letter"
    for line in body.split("\n"):
        line = line.strip()
        if line.startswith("### "):
            heading = line[4:].strip()
            continue
        if not line or line.startswith(("**To:**", "**Re:**", "*Draft")):
            continue
        sents = []
        for chunk in re.split(r"(?<=\])\s+(?=[A-Z“\"])", line):
            ids = [numbered[int(n) - 1].id for n in CITE.findall(chunk) if 1 <= int(n) <= len(numbered)]
            clean = CITE.sub("", chunk).strip()
            if clean:
                sents.append({"text": clean, "sources": ids})
        if sents:
            if sections and sections[-1]["heading"] == heading:
                sections[-1]["sentences"].extend(sents)
            else:
                sections.append({"heading": heading, "sentences": sents})
    return {"sections": sections}


def to_docx(text: str, title: str, note: str) -> bytes:
    """A Word version of the letter (markdown headings, bold labels and the numbered source list)."""
    import io

    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.core_properties.title = title
    doc.styles["Normal"].font.size = Pt(11)

    def add_runs(par, line):
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", line)):
            if part:
                run = par.add_run(part)
                run.bold = i % 2 == 1

    for line in text.split("\n"):
        line = line.rstrip()
        if not line:
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif re.match(r"^\*[^*].*\*$", line):
            doc.add_paragraph().add_run(line.strip("*")).italic = True
        else:
            add_runs(doc.add_paragraph(), line)
    footer = doc.add_paragraph()
    footer.add_run(note).italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render(draft: dict, sources: list[Source], case_name: str) -> str:
    """Markdown letter with numbered citations and a source list."""
    order: list[str] = [s.id for s in numbered_sources(draft, sources)]
    num = {i: n for n, i in enumerate(order, 1)}
    by_id = {s.id: s for s in sources}
    out = [f"**To:** {RECIPIENT}", f"**Re:** {case_name}", "",
           "*Draft for lawyer review. Not to be sent until approved.*", ""]
    for sec in draft["sections"]:
        out.append(f"### {sec['heading']}")
        out.append(" ".join(
            f"{s['text']} " + "".join(f"[{num[i]}]" for i in s["sources"] if i in num) for s in sec["sentences"]
        ))
        out.append("")
    out.append("### Sources")
    for i in order:
        s = by_id.get(i)
        if s:
            tag = "" if s.verified else " (unverified)"
            out.append(f"{num[i]}. {s.label}{tag}" + (f" — {s.url}" if s.url else ""))
    return "\n".join(out)


__all__ = ["LLMError", "SensitiveCaseError", "build_sources", "check_edited", "check_support", "generate_draft",
           "numbered_sources", "parse_edited", "render", "sentences_of", "template_draft", "to_docx", "validate"]
